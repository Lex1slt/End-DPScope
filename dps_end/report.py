"""Excel报表生成。

主表「计算过程」：顶部为总结性数据，下方按时间顺序每一行列出一段伤害，
右侧依次给出 实际伤害(非暴击)/暴击伤害（可与游戏实战逐hit对照）、
各大乘区数值与其来源明细（完整的伤害推导过程）。
辅助表：总览 / 技能与元素分布 / SP与能量 / 增益覆盖统计。
数据来自 Endaxis 模拟器（vendor/endaxis）的输出。
"""

from __future__ import annotations

import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .akedata import load_names

INK = "1A1A1C"  # 终末地墨黑
YELLOW = "FFD100"  # 结束剂黄
TITLE_FONT = Font(bold=True, size=15, color=YELLOW)
SECTION_FONT = Font(bold=True, size=12, color=INK)
HEADER_FILL = PatternFill("solid", fgColor=INK)
TOTAL_FILL = PatternFill("solid", fgColor="FFF4CC")
INFO_FILL = PatternFill("solid", fgColor="FFF4CC")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
BOLD = Font(bold=True)
THIN = Border(*[Side(style="thin", color="D9DCE0")] * 4)
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(wrap_text=True, vertical="top")


def hazard_bar(ws, row: int, col_start: int, col_end: int):
    """黑黄相间的警戒条纹细条（终末地工业风装饰）。"""
    ws.row_dimensions[row].height = 3.5
    yellow = PatternFill("solid", fgColor=YELLOW)
    ink = PatternFill("solid", fgColor=INK)
    for c in range(col_start, col_end + 1):
        ws.cell(row=row, column=c).fill = yellow if (c - col_start) % 2 == 0 else ink

ELEMENT_NAMES = {"physical": "物理", "heat": "灼热", "cryo": "寒冷",
                 "electric": "电磁", "nature": "自然", "ether": "超域"}
SKILL_TYPE_NAMES = {"basicAttack": "普攻", "battleSkill": "战技", "comboSkill": "连携技",
                    "ultimate": "终结技", "finisher": "处决", "dive": "下落攻击",
                    "finalStrike": "重击", "dot": "持续伤害", "reaction": "反应伤害",
                    "other": "其他"}
REACTION_NAMES = {"artsBurst": "法术爆发", "corrosion": "腐蚀", "electrification": "导电",
                  "combustion": "燃烧", "combustion_dot": "燃烧(持续)", "freeze": "冻结",
                  "shatter": "碎冰"}
TRIGGER_LABELS = {
    "lastrite-mirages-hit": "幻影斩",
    "lastrite-hypothermic-perfusion": "幻影斩",
    "arcane-comboSkill-trigger4-effect0": "禁锢破碎",
    "arcane-ultimate-trigger0-effect0": "落点打击",
    "arcane-ultimate-trigger0-effect1": "落点打击",
    "arcane-ultimate-trigger0-effect2": "落点打击",
    "arcane-ultimate-trigger0-effect3": "落点打击",
    "arcane-ultimate-trigger1-effect0": "落点打击",
    "arcane-ultimate-trigger1-effect1": "落点打击",
    "arcane-ultimate-trigger1-effect2": "落点打击",
    "arcane-ultimate-trigger1-effect3": "落点打击",
    "tangtang-ultimate-trigger0-effect0": "凝视爆发",
    "tangtang-ultimate-early-end-hit": "凝视提前终结",
}


# Endaxis 数据表内效果名为英文，此处映射为游戏中文名（可按需扩充）
# Endaxis 数据表内效果名为英文，此处映射为游戏中文名。
# 名称以官方中文语料（AKE_I18nTextTable_CN）逐一核对：技能/天赋/武器/套装用官方名，
# 增益用描述文本里出现的官方称呼（如 低温灌注/囹圄/支援晶体/集束打击）。
LABEL_CN = {
    # 武器
    "Khravengger": "赫拉芬格",
    "Brigand's Calling": "落草",
    "Type 42: Solemn Phalanx": "四二式·肃阵",
    "Chivalric Virtues": "骑士精神",
    # 装备件 / 套装
    "Tide Surge Gauntlets": "潮涌手甲",
    "Hanging River O2 Tube": "悬河供氧栓",
    "Tide Surge": "潮涌",
    "Eternal Xiranite": "长息",
    "Qingbo": "清波",
    "Bonekrusha": "碾骨",
    "Bonekrusha Heavy Armor T2": "碾骨重护甲·贰型",
    "Bonekrusha Wristband T1": "碾骨腕带·壹型",
    "Bonekrusha Mask T2": "碾骨面甲·贰型",
    # 干员技能 / 天赋（官方名）
    "Hypothermia": "低温症",
    "Tactical Planning": "筹谋",
    "Riot Bringer": "呼风唤浪",
    "bonekrushingSmash": "碾骨重压",
    # 增益 / 减益（官方描述里的称呼）
    "auxiliaryCrystal": "支援晶体",
    "xaihi-auxiliary-crystal": "支援晶体",
    "xaihi-auxiliary-crystal-amp": "支援晶体",
    "xaihi-ultimate-cryo-amp": "寒冷增幅",
    "xaihi-ultimate-nature-amp": "自然增幅",
    "waterspouts": "水龙卷",
    "tangtang-t2-waterspouts": "水龙卷",
    "oldenStare": "凝视",
    "tangtang-oldenStare": "凝视",
    "tangtang-ultimate-trigger0-effect0": "大当家盯着呢！",
    "whirlpools": "涡流",
    "lastrite-hypothermic-perfusion": "低温灌注",
    "lastrite-mirages-hit": "低温灌注",
    "arcane-imprisonment": "囹圄",
    "arcane-comboSkill-trigger4-effect0": "囹圄",
    "arcane-ultimate-trigger0-effect0": "集束打击",
    "arcane-ultimate-trigger0-effect1": "集束打击",
    "arcane-ultimate-trigger0-effect2": "集束打击",
    "arcane-ultimate-trigger0-effect3": "集束打击",
    "arcane-ultimate-trigger1-effect0": "集束打击",
    "arcane-ultimate-trigger1-effect1": "集束打击",
    "arcane-ultimate-trigger1-effect2": "集束打击",
    "arcane-ultimate-trigger1-effect3": "集束打击",
    "arcane-ultimate-cluster-strike-counter": "集束打击（计数）",
    "arcane-gloompurger-array": "破晦阵",
    "arcane-gloompurge-arcana-ready": "破晦诀明",
    "arcane-combo-susceptibility": "自然/寒冷脆弱",
    "natureCryoSusceptibility": "自然/寒冷脆弱",
    # 反应 / 通用键
    "corrosion": "腐蚀",
    "corrosion:resShred": "腐蚀",
    "resShred": "腐蚀",
    "susceptibility:arts": "法术脆弱",
    "increasedDmgTaken:cryo": "寒冷脆弱",
    "ampBonus:cryo": "寒冷增幅",
    "ampBonus:nature": "自然增幅",
    "dmgBonus": "伤害加成",
    # 敌方减益实例ID → 中文名
    "type-42-solemn-phalanx-increasedDmgTaken-susceptibility": "四二式·肃阵",
    "type-42-solemn-phalanx-increasedDmgTaken-burst": "四二式·肃阵（爆发）",
    "last-rite-talent0-trigger0-effect0": "低温症",
    "arcane-t1": "筹谋",
    "tangtang-waterspouts-susceptibility": "呼风唤浪",
    "brigands-calling-skill3-trigger0-effect0": "落草",
    "brigands-calling-skill3-trigger1-effect0": "落草",
    "khravengger-skill3-trigger0-effect0": "赫拉芬格",
    "khravengger-skill3-trigger1-effect0": "赫拉芬格",
    "chivalric-virtues-skill3-trigger0-effect0": "骑士精神",
    "qingbo-trigger0-effect0": "清波",
    "bonekrusha-trigger0-effect0": "碾骨重压",
    "tide-surge-trigger0-effect0": "潮涌",
    "eternal-xiranite-set": "长息",
    "staggered": "失衡",
    "last-rite-combo-window": "别礼·连携窗口",
    "tangtang-combo-window": "汤汤·连携窗口",
    "xaihi-combo-window": "赛希·连携窗口",
    "arcane-combo-window": "诀·连携窗口",
    "arcane-combo-tracker-cryo": "诀·连携计数（寒冷）",
    "arcane-combo-tracker-nature": "诀·连携计数（自然）",
    "arcane-combo-tracker-extender-cryo": "诀·连携计数延长（寒冷）",
    "tangtang-whirlpools": "汤汤·涡流",
    # 提弗洛斯（typhoeus，上游 d248cf30 新增）——官方名均经 AKE_I18nTextTable_CN 子串检索核验
    "Samifjod Fletching": "萨米制箭工艺",
    "Umbra of Frigid Eventide": "寒夜幽影",
    "typhoeus-sign": "启示",
    "typhoeus-hunting-arrow": "猎矢",
    "typhoeus-hovering": "浮空",
    "typhoeus-hail-of-arrows": "箭雨",
    "typhoeus-barrage-array": "箭阵",
    "typhoeus-barrage-array-burst-dmg-taken": "箭阵",
    "typhoeus-talent0-effect0": "猎物清点",
    "typhoeus-potential0-effect0": "萨米制箭工艺",
    "typhoeus-potential0-effect1": "萨米制箭工艺",
    "windform": "身形如风",
    "umbra-of-frigid-skill3-trigger0-effect0": "寒夜幽影",
    "slowed": "缓速",
}

# 增益「类别」列里的属性修饰符 → 中文
STAT_NAMES = {
    "dmgBonus": "伤害加成", "ampBonus": "增幅", "atkPercent": "攻击力", "atkFlat": "攻击力",
    "critRate": "暴击率", "critDmg": "暴击伤害", "hpPercent": "生命", "flatHp": "生命",
    "attributeFlat": "属性", "attributePercent": "属性", "susceptibility": "易伤",
    "increasedDmgTaken": "受伤提升", "ultimateGainEfficiency": "终结技充能效率",
    "ultimateEnergyCostReduction": "终结技充能效率", "protection": "防护",
    "artsIntensity": "法术强度", "heal": "治疗", "resShred": "抗性削减",
    "cryoInfliction": "寒冷附着", "natureInfliction": "自然附着",
    "heatInfliction": "灼热附着", "electricInfliction": "电磁附着",
    "physicalInfliction": "物理附着",
}

# SP 事件流水里的 reason → 中文
SP_REASON_NAMES = {
    "skill": "技能消耗", "damage": "造成伤害回复", "hit": "命中回复",
    "finisher": "处决回复", "regen": "自然恢复", "extra": "额外回复",
}

# 效果ID语素兜底翻译：形如「干员-效果名」的复合ID按语素拼接
MORPHEMES = {
    "combo": "连携", "window": "窗口", "tracker": "计数", "extender": "延长",
    "trigger": "触发", "set": "套装", "susceptibility": "易伤", "auxiliary": "辅助",
    "crystal": "晶体", "amp": "增幅", "talent": "天赋", "skill": "技能",
    "burst": "爆发", "whirlpools": "涡流", "waterspouts": "水龙卷", "oldenStare": "凝视",
    "gloompurger": "弥尘", "array": "法阵", "arcana": "法术", "ready": "就绪",
    "hypothermic": "低温", "perfusion": "灌注", "mirages": "幻影", "hit": "斩",
    "increasedDmgTaken": "受伤提高", "resShred": "腐蚀", "staggered": "失衡",
    "comboSkill": "连携技", "ultimate": "终结技", "battleSkill": "战技",
    "basicAttack": "普攻", "finalStrike": "重击", "counter": "计数",
    "cryo": "寒冷", "nature": "自然", "heat": "灼热", "electric": "电磁", "physical": "物理",
    "arcane": "诀", "lastrite": "别礼", "tangtang": "汤汤", "xaihi": "赛希",
    "qingbo": "清波", "bonekrusha": "碾骨", "eternal": "长息", "xiranite": "长息",
    "chivalric": "骑士", "virtues": "精神", "khravengger": "赫拉芬格",
    "brigands": "落草", "calling": "落草", "solemn": "肃阵", "phalanx": "肃阵",
    "tide": "潮涌", "surge": "潮涌", "hanging": "悬河", "river": "悬河", "type-42": "四二式",
}


def _morpheme_zh(label: str) -> str | None:
    """把「干员-效果名」式英文ID按语素拼成中文；拼不出来的返回 None。"""
    parts = label.split("-")
    if len(parts) < 2 or any(not p for p in parts):
        return None
    if not all(p in MORPHEMES for p in parts):
        return None
    out = "".join(MORPHEMES[p] for p in parts)
    return out if out != label else None


def T(label: str | None) -> str:
    if not label:
        return "?"
    if label in LABEL_CN:
        return LABEL_CN[label]
    if label in STAT_NAMES:
        return STAT_NAMES[label]
    m = re.fullmatch(r"(physical|heat|cryo|electric|nature)Infliction", label)
    if m:
        return f"{ELEMENT_NAMES[m.group(1)]}附着"
    if label in ELEMENT_NAMES:
        return ELEMENT_NAMES[label]
    return _morpheme_zh(label) or label


def _pct(v: float, signed: bool = True, digits: int = 2) -> str:
    s = f"{v * 100:.{digits}f}%"
    if signed:
        return ("+" if v >= 0 else "−") + s.lstrip("+")
    return s


def _style_header(ws, row, cols, underline=YELLOW):
    """黑色表头带 + 彩色下划线（默认终末地黄）。"""
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = Border(
            bottom=Side(style="medium", color=underline),
            top=Side(style="thin", color=INK),
            left=Side(style="thin", color=INK),
            right=Side(style="thin", color=INK),
        )


LMDI_METHOD_LINES = [
    ("四、计算方法（LMDI · 供核对与讨论）", True),
    ("1. 基线：对每一击，保留命中者自身增益与基础面板，剔除全部队友提供的增益/减益/连击/失衡贡献，"
     "用同一套伤害公式重算一次，得到\"无拐基线伤害\" D_self；该击实际期望伤害记为 D_actual。", False),
    ("2. 若 D_actual ≤ D_self（外部贡献为零或为负），该击全部记为自身。", False),
    ("3. 期望伤害 = 基础 × (1+加成) × (1+增幅) × 暴击期望 × (1+易伤) × (1+承伤) × 抗性 × 连击 × 失衡 × …。"
     "对每个乘区因子 k，其贡献 = L × ln(f_actual / f_self)，其中对数平均 L(a,b) = (a−b)/ln(a/b)；"
     "乘区之间的交互影响按对数对称分摊，这是乘法分解的标准 LMDI 方法（对数平均迪氏指数）。", False),
    ("4. 同一乘区内有多个来源时（例如易伤 +20% = 甲的 +15% + 乙的 +5%），按各自加成数值占比分摊（75% : 25%）。", False),
    ("5. 自身 = D_actual − Σ 外部贡献（余项）。因此\"自身 + Σ拐力 = 该击期望\"逐击 精确成立，"
     "全队合计恰等于总伤害，可在本表逐行核对。", False),
    ("口径说明：", True),
    ("- 记账采用\"施加者\"模式：敌人身上的易伤/承伤/抗性削减，整笔记给施加该减益的干员；"
     "连击层数与失衡值按各干员的贡献占比拆分。", False),
    ("- 分解对象是期望伤害（暴击期望 = 1 + 暴击率 × 暴击伤害），不含单次暴击掷点的运气成分。", False),
    ("- 防御区、抗性区基线不受干员增益影响，不产生拐力；反应伤害（法术爆发/燃烧/冻结等）"
     "按消耗的附着层数单独记功。", False),
    ("- 基线是\"该击剔除外部增益后重算\"，不是\"把该干员移出队伍重打一遍\"；"
     "出招序列由排轴给定，队友影响出招本身的部分不在归因范围内。", False),
    ("- 归因方法存在其他合理选择（逐个剔除法、Shapley 值等），数值会略有差异；"
     "本报告固定使用 LMDI 以保证全队精确可加、逐击可核对。", False),
]


def build_workbook(sim: dict, project: dict, names: dict | None = None) -> Workbook:
    names = names or load_names()
    op_names = names.get("operators", {})

    def dn(track_id: str | None) -> str:
        return op_names.get(track_id or "", track_id or "")

    wb = Workbook()
    hits = [h for h in sim.get("hits", []) if h.get("exp") is not None]
    damaging = [h for h in hits if (h.get("exp") or 0) > 0]
    for h in damaging:
        if not h.get("st"):
            h["st"] = "reaction" if h.get("reaction") else "other"
    summary = sim.get("summary", {})
    total = summary.get("totalDamage", 0)
    total_nc = sum(h.get("nc") or 0 for h in damaging)
    total_cc = sum(h.get("crit") or 0 for h in damaging)
    dps = summary.get("dps") or 0

    tracks = (project.get("scenarioList") or [{}])[0].get("data", {}).get("tracks", [])
    track_ids = [t.get("id") for t in tracks]
    sp_cost = sum(-s["sp"] for s in sim.get("spLog", [])
                  if s.get("kind") == "SP_CHANGE" and (s.get("sp") or 0) < 0)
    dpsp = total / sp_cost * 100 if sp_cost else 0  # 每格技力
    enemy_name = (names.get("enemies", {}).get(sim.get("meta", {}).get("activeEnemyId", ""))
                  or sim.get("meta", {}).get("enemyName") or "未知敌人")
    enemy_hp_left = summary.get("enemyHpLeft")

    # ── 标注方案：动作列 = "汤汤战技hit1"式完整标签 ──
    # 动作hit序号：同一动作内的伤害命中依次 hit1/hit2/...
    act_damage_ordinal: dict = defaultdict(int)
    for e in sorted(damaging, key=lambda x: (x["t"], x.get("src") or "")):
        act = e.get("act")
        if act and not (e.get("triggeredBy") or "").startswith("dot:") and not e.get("triggeredBy"):
            act_damage_ordinal[act] += 1
            e["_hit_no"] = act_damage_ordinal[act]
    # DoT跳伤序号：按DoT实例分组
    dot_ordinal: dict = defaultdict(int)
    for e in sorted(damaging, key=lambda x: x["t"]):
        trig = e.get("triggeredBy") or ""
        if trig.startswith("dot:"):
            dot_ordinal[trig] += 1
            e["_dot_no"] = dot_ordinal[trig]
    # 触发hit序号：按(动作,触发来源)分组
    trig_ordinal: dict = defaultdict(int)
    for e in sorted(damaging, key=lambda x: (x["t"], x.get("src") or "")):
        trig = e.get("triggeredBy") or ""
        if trig and not trig.startswith("dot:"):
            key = (e.get("act"), trig)
            trig_ordinal[key] += 1
            e["_trig_no"] = trig_ordinal[key]

    def hit_label(e: dict) -> str:
        op = dn(e.get("src") or "")
        trig = e.get("triggeredBy") or ""
        if h_get_reaction := e.get("reaction"):
            base = REACTION_NAMES.get(h_get_reaction, h_get_reaction)
            return f"{op}{base}"
        if trig.startswith("dot:"):
            dot = trig[4:].split("@")[0]
            dot = LABEL_CN.get(dot) or TRIGGER_LABELS.get(dot, dot)
            return f"{op}{dot}跳伤{e.get('_dot_no', '')}"
        for key, label in TRIGGER_LABELS.items():
            if key in trig or trig == key:
                n = e.get("_trig_no")
                return f"{op}{label}" + (f"第{n}段" if (n or 1) > 1 or key != trig else "")
        if e.get("st") == "finisher":
            return f"{op}处决"
        action = e.get("actionName") or "伤害"
        n = e.get("_hit_no")
        return f"{op}{action}" + (f"第{n}段" if n else "")

    # ══════════════ Sheet 1 计算过程（主表） ══════════════
    ws = wb.active
    ws.title = "计算过程"
    ws.sheet_view.showGridLines = False
    col_widths = [2, 8, 8, 18, 8, 9, 30, 10,
                  11, 11, 11, 9, 9, 11, 30, 9,
                  11, 8, 28, 8, 28, 8, 24, 8, 24,
                  8, 8, 24, 8, 12, 12, 24]
    for i, w in enumerate(col_widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = w

    ACCENT_BAR = PatternFill("solid", fgColor=INK)
    METRIC_FILL = PatternFill("solid", fgColor="FFF6D9")
    METRIC_LABEL_FONT = Font(size=8, color="8A8F98")

    # ---- 顶部总结区 ----
    r = 2
    for col in range(2, 15):
        ws.cell(r, col).fill = ACCENT_BAR
    ws.cell(r, 2, "End-DPScope · 终末地排轴伤害计算过程").font = TITLE_FONT
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=14)
    ws.row_dimensions[r].height = 26
    r += 1
    hazard_bar(ws, r, 2, 14)
    r += 2

    if enemy_hp_left is None:
        left_txt, left_color = "—", "8496B0"
    elif enemy_hp_left <= 0:
        left_txt, left_color = f"已击杀（超出 {-enemy_hp_left:,.0f}）", "0FA968"
    else:
        left_txt, left_color = f"{enemy_hp_left:,.0f}", "C00000"
    metrics = [
        ("总伤害（期望）", f"{total:,.0f}", "1F3864"),
        ("DPS", f"{dps:,.1f}", "1F3864"),
        ("DPSP", f"{dpsp:,.1f}", "1F3864"),
        ("敌人剩余血量", left_txt, left_color),
    ]
    label_row, value_row = r, r + 1
    for j, (lab, val, color) in enumerate(metrics):
        c1 = 2 + j * 3
        ws.merge_cells(start_row=label_row, start_column=c1, end_row=label_row, end_column=c1 + 2)
        ws.merge_cells(start_row=value_row, start_column=c1, end_row=value_row, end_column=c1 + 2)
        lc = ws.cell(label_row, c1, lab)
        lc.font = METRIC_LABEL_FONT
        lc.alignment = CENTER
        vc = ws.cell(value_row, c1, val)
        vc.font = Font(bold=True, size=14, color=color)
        vc.alignment = CENTER
        for k in range(3):
            ws.cell(value_row, c1 + k).fill = METRIC_FILL
    ws.row_dimensions[value_row].height = 22
    r = value_row + 2

    info_rows = [
        ("方案 / 来源", f"{sim.get('meta', {}).get('scenario', '')} / "
                        f"{Path(sim.get('meta', {}).get('source', '') or '.').name}"),
        ("敌人 / 血量剩余", f"{enemy_name} / "
                            + (f"剩余 {enemy_hp_left:,.0f}" if enemy_hp_left is not None else "-")),
        ("统计窗口", f"{summary.get('startlineSec', 0):.2f}s → "
                     f"{summary.get('lastDamageTime', 0):.2f}s"
                     f"（截止最后一次造成伤害，循环 {summary.get('rotationTime', 0):.2f}s）"),
        ("不暴击 / 全暴击", f"{total_nc:,.0f} / {total_cc:,.0f}"
                            f"（总SP消耗 {sp_cost:,.0f}）"),
        ("计算核心", "Endaxis 模拟器（乘区/状态/触发器/法术异常/失衡/处决与网页版一致）"),
    ]
    for label, value in info_rows:
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        c = ws.cell(r, 2, label)
        c.font = BOLD
        c.fill = INFO_FILL
        c.border = THIN
        ws.cell(r, 3).border = THIN
        ws.cell(r, 3).fill = INFO_FILL
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=14)
        v = ws.cell(r, 4, value)
        v.font = Font(size=9)
        v.border = THIN
        r += 1
    r += 1

    # 干员贡献一行
    lmdi = {row["track"]: row for row in summary.get("lmdi", [])}
    parts = []
    for tid in track_ids:
        row = lmdi.get(tid, {"dmg": 0, "buff": 0})
        parts.append(f"{dn(tid)} {row.get('dmg', 0):,.0f}+拐{row.get('buff', 0):,.0f}")
    c = ws.cell(r, 2, "干员（结算+拐力）")
    c.font = BOLD
    c.fill = INFO_FILL
    c.border = THIN
    ws.cell(r, 3, "；".join(parts)).border = THIN
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
    r += 2

    # ---- 逐hit计算表 ----
    headers = ["时间(s)", "干员", "动作", "类型",
               "敌人附着", "敌人减益", "失衡",
               "期望伤害", "非暴击伤害", "暴击伤害",
               "实时暴击率", "实时暴击伤害",
               "有效攻击力", "攻击力来源", "倍率(%)", "基础伤害",
               "加成区", "加成区来源", "增幅区", "增幅区来源",
               "脆弱区", "脆弱区来源", "易伤区", "易伤区来源",
               "防御区", "抗性区", "抗性区来源", "失衡区", "处决/其他", "反应加成",
               "LMDI拐力归因"]
    # 分组配色：黑色表头 + 每组一条彩色下划线（终末地工业风）
    GR_BASE, GR_ENEMY, GR_DMG = "FFFFFF", "B39DDB", YELLOW
    GR_PANEL, GR_ZONE, GR_SRC = "AEB8C4", "7FA8D9", "9DC3E6"
    GR_MISC, GR_LMDI = "9AA0A6", "6FBF73"
    group_underline = ([GR_BASE] * 4 + [GR_ENEMY] * 3 + [GR_DMG] * 3 + [GR_PANEL] * 6
                       + [GR_ZONE, GR_SRC] * 4 + [GR_ZONE] * 3 + [GR_ZONE] + [GR_MISC] * 2
                       + [GR_LMDI])
    header_row = r
    ws.row_dimensions[header_row].height = 30
    for i, h in enumerate(headers):
        c = ws.cell(header_row, 2 + i, h)
        c.fill = HEADER_FILL
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.alignment = CENTER
        c.border = Border(bottom=Side(style="medium", color=group_underline[i]),
                          top=Side(style="thin", color=INK),
                          left=Side(style="thin", color=INK),
                          right=Side(style="thin", color=INK))
    ws.freeze_panes = ws.cell(row=header_row + 1, column=9).coordinate
    ws.auto_filter.ref = f"B{header_row}:AF{header_row}"

    zone_fmt = "0.000"

    def zone_sources_text(sources, kind="buff") -> str:
        if not sources:
            return "—" if kind != "shred" else "无削减"
        bits = []
        for s in sources:
            label = T(s.get("label"))
            v = s.get("value") or 0
            if kind == "shred":
                bits.append(f"{label} −{_pct(v, signed=False)}")
            else:
                bits.append(f"{label} {_pct(v)}")
        return "；".join(bits)

    zebra = PatternFill("solid", fgColor="F5F7FA")
    hair = Border(bottom=Side(style="thin", color="E3E6EA"))
    data_font = Font(size=9)
    data_font_bold = Font(size=9, bold=True, color="1F3864")
    dim_font = Font(size=9, color="A6A6A6")
    exp_font = Font(size=9, bold=True, color="7F6000")
    op_fills = {}
    for i, tid in enumerate(track_ids):
        if tid:
            op_fills[tid] = PatternFill(
                "solid", fgColor=["DDEBF7", "E2EFDA", "FFF2CC", "FCE4D6", "E4DFEC"][i % 5])
    text_cols = {2, 5, 13, 17, 19, 21, 23, 26, 28, 29, 30}

    z = header_row + 1
    for e in sorted(damaging, key=lambda x: (x["t"], x.get("src") or "")):
        ext = e.get("lmdiExt") or {}
        exp_v = e.get("exp") or 0
        ext_sum = sum(abs(v) for v in ext.values())
        attr_bits = []
        if ext_sum > 0.5 and exp_v:
            attr_bits.append(f"自身 {1 - ext_sum / exp_v:.0%}")
        attr_bits += (f"{dn(s)} {v:,.0f}（{v / exp_v:.0%}）" for s, v in
                      sorted(ext.items(), key=lambda kv: -kv[1]) if abs(v) > 0.5)
        attr_txt = "；".join(attr_bits) if attr_bits else "无外部增益"

        # 攻击力来源
        atk_bits = []
        for s in e.get("atkSources") or []:
            label = T(s.get("label"))
            v = s.get("value") or 0
            if label == "__base__":
                continue
            atk_bits.append(f"{label} {_pct(v)}")
        flat = e.get("flatAtk") or 0
        if flat:
            atk_bits.append(f"固定+{flat:g}")
        atk_text = ("、".join(atk_bits) + f" → {e.get('atk'):,.0f}") if atk_bits else \
                   f"面板 {e.get('atk'):,.0f}"

        # 抗性区来源
        res_bits = []
        enemy_res = e.get("enemyRes")
        if enemy_res:
            res_bits.append(f"基础抗性{_pct(enemy_res, signed=False)}")
        if e.get("resIgnore"):
            res_bits.append(f"无视 {_pct(e['resIgnore'], signed=False)}")
        shred_txt = zone_sources_text(e.get("shredSources"), kind="shred")
        if shred_txt != "无削减":
            res_bits.append(f"削减：{shred_txt}")
        res_text = "；".join(res_bits) if res_bits else "无抗性"

        # 处决/其他
        misc = []
        if (e.get("fin") or 1) != 1:
            misc.append(f"处决承伤×{e['fin']:g}")
        if (e.get("direct") or 1) != 1:
            misc.append(f"直接倍率×{e['direct']:g}")
        if (e.get("linkStacks") or 0) > 0:
            misc.append(f"连击{e['linkStacks']}层×{e.get('link'):.2f}")
        misc_text = "；".join(misc) if misc else "—"

        # 反应加成
        if e.get("reaction"):
            bits = []
            if e.get("reactionLevel"):
                bits.append(f"等级系数×{e['reactionLevel']:.3f}")
            if e.get("reactionArts"):
                bits.append(f"源石技艺×{e['reactionArts']:.3f}")
            eff = e.get("reactionEffect")
            if eff and eff != 1:
                bits.append(f"效果强化×{eff:g}")
            reaction_text = " ".join(bits) if bits else "—"
        else:
            reaction_text = "—"

        # 实时暴击面板（分数；伊冯等叠暴机制下逐击 变化）
        cr, cd = e.get("critRate"), e.get("critDmg")

        es = e.get("enemyState") or {}
        attach_txt = "—"
        inf = es.get("infliction")
        if inf:
            attach_txt = f"{ELEMENT_NAMES.get(inf.get('element'), inf.get('element'))}×{inf.get('stacks')}"
        debuff_bits = []
        seen_debuffs: dict = defaultdict(float)
        seen_stats: dict = {}
        for dbf in es.get("debuffs") or []:
            did = dbf.get("id") or "?"
            seen_debuffs[did] += dbf.get("value") or 0
            if dbf.get("stat"):
                seen_stats[did] = dbf["stat"]
        for did, v in seen_debuffs.items():
            # 快照里无显式 id 的减益（引擎随机 uid）用 stat 修饰符兜底取官方名
            label = LABEL_CN.get(did) or LABEL_CN.get(seen_stats.get(did, "")) or did
            if abs(v) > 0.01:
                debuff_bits.append(f"{label} {_pct(v / 100, signed=False)}")
            else:
                debuff_bits.append(label)
        debuff_txt = "；".join(debuff_bits) if debuff_bits else "—"
        stagger_txt = "失衡中" if es.get("broken") else (
            f"{es.get('stagger', 0):g}/{es.get('maxStagger', 0):g}" if es.get("maxStagger") else "—")

        vals = [
            e["t"], dn(e.get("src") or ""),
            hit_label(e),
            SKILL_TYPE_NAMES.get(e.get("st") or "", e.get("st") or ""),
            attach_txt, debuff_txt, stagger_txt,
            e.get("exp"), e.get("nc"), e.get("crit"),
            cr, cd,
            e.get("atk"), atk_text,
            e.get("resolvedMult") if e.get("resolvedMult") is not None else e.get("mult"),
            e.get("base"),
            e.get("bonus"), zone_sources_text(e.get("dmgSources")),
            e.get("amp"), zone_sources_text(e.get("ampSources")),
            e.get("susc"), zone_sources_text(e.get("suscSources")),
            e.get("taken"), zone_sources_text(e.get("takenSources")),
            e.get("def"), e.get("res"), res_text,
            e.get("stagger"), misc_text, reaction_text,
            attr_txt,
        ]
        op_fill = op_fills.get(e.get("src") or "")
        for i, v in enumerate(vals):
            c = ws.cell(z, 2 + i, v)
            c.font = data_font
            c.border = hair
            if i == 1:
                if op_fill:
                    c.fill = op_fill
                    c.font = data_font_bold
            elif z % 2 == 0:
                c.fill = zebra
            if i == 0:
                c.number_format = "0.00"
            elif i in (8, 9, 12, 15):
                c.number_format = "#,##0"
                c.font = data_font_bold
            elif i == 7:
                c.number_format = "#,##0"
                c.font = exp_font
            elif i in (10, 11):
                c.number_format = "0.0%"
                if v is not None:
                    c.font = data_font
            elif i in (16, 18, 20, 22, 24, 25, 27):
                c.number_format = zone_fmt
            elif i == 14:
                c.number_format = "0.##"
            if v == "—":
                c.font = dim_font
            if i in text_cols:
                c.alignment = WRAP
            else:
                c.alignment = CENTER
        ws.row_dimensions[z].height = 26
        z += 1

    # ══════════════ Sheet 2 总览 ══════════════
    ws0 = wb.create_sheet("总览")
    ws0.sheet_view.showGridLines = False
    ws0.column_dimensions["A"].width = 2
    for col, w in zip("BCDEFGHIJKLM", (24, 15, 14, 14, 14, 14, 14, 12, 12, 12, 12, 12)):
        ws0.column_dimensions[col].width = w
    r0 = 2
    for col in range(2, 14):
        ws0.cell(r0, col).fill = PatternFill("solid", fgColor=INK)
    ws0.cell(r0, 2, "End-DPScope · 终末地排轴伤害报告").font = TITLE_FONT
    ws0.merge_cells(start_row=r0, start_column=2, end_row=r0, end_column=13)
    ws0.row_dimensions[r0].height = 26
    r0 += 1
    hazard_bar(ws0, r0, 2, 13)
    r0 += 2
    for k, v in [
        ("方案 / 来源", f"{sim.get('meta', {}).get('scenario', '')} / "
                        f"{Path(sim.get('meta', {}).get('source', '') or '.').name}"),
        ("敌人", enemy_name),
        ("敌人HP / 剩余", (f"{total + enemy_hp_left:,.0f} / {enemy_hp_left:,.0f}"
                           if enemy_hp_left is not None else "-")),
        ("统计窗口", f"至最后一次造成伤害 {summary.get('lastDamageTime', 0):.2f}s"
                    f"（起点 {summary.get('startlineSec', 0):.2f}s）"),
    ]:
        ws0.cell(r0, 2, k).font = BOLD
        ws0.cell(r0, 3, v)
        r0 += 1
    ec = summary.get("enemyConfig") or {}
    if ec:
        res = ec.get("resistance") or {}
        res_txt = " / ".join(f"{ELEMENT_NAMES.get(k, k)} {v:g}" for k, v in res.items()) or "—"
        for k, v in [
            ("本次敌人配置", f"HP {ec.get('hp') or 0:,.0f} · 防御 {ec.get('defense', '—')} · "
                            f"失衡上限 {ec.get('maxStagger', '—')} · "
                            f"失衡时长 {ec.get('staggerBreakDuration', '—')}s · "
                            f"处决回技 {ec.get('executionRecovery', '—')}"),
            ("敌人抗性", res_txt),
            ("面板修正条数", ec.get("panelOverrides", 0)),
        ]:
            ws0.cell(r0, 2, k).font = BOLD
            ws0.cell(r0, 3, v)
            r0 += 1
    r0 += 1
    totals = [
        ("总伤害（期望）", total, "#,##0"),
        ("DPS（伤害/秒）", dps, "#,##0.0"),
        ("总SP消耗", sp_cost, "#,##0"),
        ("DPSP", dpsp, "#,##0.0"),
        ("总伤害（不暴击）", total_nc, "#,##0"),
        ("总伤害（全暴击）", total_cc, "#,##0"),
    ]
    for i, (k, v, fmt) in enumerate(totals):
        col = 2 + (i % 3) * 2
        row = r0 + i // 3
        ws0.cell(row, col, k).font = BOLD
        c = ws0.cell(row, col + 1, v)
        c.number_format = fmt
        c.fill = TOTAL_FILL
    r0 += (len(totals) + 2) // 3 + 1

    ws0.cell(r0, 2, "干员贡献（LMDI归因）").font = SECTION_FONT
    r0 += 1
    headers0 = ["干员", "结算伤害", "伤害占比", "拐力（增益贡献）", "拐力占比",
                "实际贡献", "实际贡献占比", "SP消耗", "终结技次数"]
    for i, h in enumerate(headers0):
        ws0.cell(r0, 2 + i, h)
    _style_header(ws0, r0, len(headers0) + 1)
    r0 += 1
    rows0 = []
    for tid in track_ids:
        row = lmdi.get(tid, {"dmg": 0.0, "buff": 0.0})
        dealt, buff = row.get("dmg", 0.0), row.get("buff", 0.0)
        actions = next((t for t in tracks if t.get("id") == tid), {}).get("actions", [])
        ults = sum(1 for a in actions if a.get("type") == "ultimate")
        rows0.append([dn(tid), dealt, dealt / total if total else 0, buff,
                      buff / total if total else 0, dealt + buff,
                      (dealt + buff) / total if total else 0, sp_by_track(tid, sim), ults])
    rows0.sort(key=lambda x: -x[5])
    for rd in rows0:
        for i, v in enumerate(rd):
            c = ws0.cell(r0, 2 + i, v)
            c.border = THIN
            if i in (1, 3, 5):
                c.number_format = "#,##0"
            elif i in (2, 4, 6):
                c.number_format = "0.00%"
            elif i == 7:
                c.number_format = "#,##0"
        r0 += 1
    ws0.cell(r0, 2, "合计").font = BOLD
    sums = [sum(row[i] for row in rows0) for i in range(1, 9)]
    for i, v in enumerate(sums):
        c = ws0.cell(r0, 3 + i, v)
        c.font = BOLD
        c.fill = TOTAL_FILL
        if i in (0, 2, 4, 6):
            c.number_format = "#,##0"
        elif i in (1, 3, 5):
            c.number_format = "0.00%"
    r0 += 2

    # ══════════════ Sheet 2.5 拐力归因（LMDI） ══════════════
    wsl = wb.create_sheet("拐力归因")
    wsl.sheet_view.showGridLines = False
    wsl.column_dimensions["A"].width = 2
    wsl.column_dimensions["B"].width = 24
    for col in "CDEFGHIJK":
        wsl.column_dimensions[col].width = 16
    rl = 2
    for col in range(2, 11):
        wsl.cell(rl, col).fill = PatternFill("solid", fgColor=INK)
    wsl.cell(rl, 2, "拐力归因（LMDI · 谁拐了谁，拐了多少）").font = TITLE_FONT
    wsl.merge_cells(start_row=rl, start_column=2, end_row=rl, end_column=10)
    wsl.row_dimensions[rl].height = 26
    rl += 1
    hazard_bar(wsl, rl, 2, 10)
    rl += 2
    wsl.cell(rl, 2, "LMDI 分解保证：全队「自身结算 + 拐力」之和恰等于总伤害，且逐击成立；"
                    "「自身结算」可理解为该干员去掉所有队友增益后的无拐基线。").font = BOLD
    wsl.cell(rl, 2).alignment = WRAP
    wsl.merge_cells(start_row=rl, start_column=2, end_row=rl + 1, end_column=10)
    rl += 3

    # 逐hit聚合：受拐与拐力去向
    received: dict[str, float] = defaultdict(float)          # 该干员命中中来自队友的部分
    matrix: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    hit_contribs = []                                         # (拐人, hit, v)
    for e in damaging:
        for s, v in (e.get("lmdiExt") or {}).items():
            if abs(v) <= 0.5:
                continue
            received[e.get("src")] += v
            matrix[s][e.get("src")] += v
            hit_contribs.append((s, e, v))

    # ---- 一、干员贡献详解 ----
    wsl.cell(rl, 2, "一、干员贡献详解").font = SECTION_FONT
    rl += 1
    headers_l = ["干员", "自身结算（无拐基线）", "无拐占比", "受拐（其命中中队友的增益）",
                 "自身命中合计", "拐力（提供给他人的增益）", "拐力占比",
                 "实际贡献（结算+拐力）", "实际贡献占比"]
    for i, h in enumerate(headers_l):
        wsl.cell(rl, 2 + i, h)
    _style_header(wsl, rl, len(headers_l) + 1)
    rl += 1
    for tid in track_ids:
        row = lmdi.get(tid, {"dmg": 0.0, "buff": 0.0})
        dealt, buff = row.get("dmg", 0.0), row.get("buff", 0.0)
        recv = received.get(tid, 0.0)
        vals = [dn(tid), dealt, dealt / total if total else 0, recv,
                dealt + recv, buff, buff / total if total else 0,
                dealt + buff, (dealt + buff) / total if total else 0]
        for i, v in enumerate(vals):
            c = wsl.cell(rl, 2 + i, v)
            c.border = THIN
            if i in (1, 3, 4, 5, 7):
                c.number_format = "#,##0"
            elif i in (2, 6, 8):
                c.number_format = "0.00%"
        rl += 1
    sums_l = [sum(lmdi.get(t, {"dmg": 0, "buff": 0}).get("dmg", 0) for t in track_ids),
              sum(received.get(t, 0.0) for t in track_ids),
              sum(lmdi.get(t, {"dmg": 0, "buff": 0}).get("buff", 0) for t in track_ids)]
    wsl.cell(rl, 2, "合计").font = BOLD
    for i, v in [(1, sums_l[0]), (3, sums_l[1]), (5, sums_l[2]), (7, total)]:
        c = wsl.cell(rl, 2 + i, v)
        c.font = BOLD
        c.fill = TOTAL_FILL
        c.number_format = "#,##0"
    c = wsl.cell(rl, 4, sums_l[0] / total if total else 0)  # 无拐占比合计
    c.font = BOLD
    c.fill = TOTAL_FILL
    c.number_format = "0.00%"
    rl += 2

    # ---- 二、拐力去向矩阵（行=拐人，列=被拐） ----
    wsl.cell(rl, 2, "二、拐力去向矩阵（行 = 增益来源干员，列 = 受益干员）").font = SECTION_FONT
    rl += 1
    benefs = sorted({b for row in matrix.values() for b in row},
                    key=lambda b: -sum(row.get(b, 0.0) for row in matrix.values()))
    benefs = [b for b in benefs if any(row.get(b, 0.0) > 0.5 for row in matrix.values())]
    if benefs:
        for i, h in enumerate(["拐人 ↓ / 被拐 →"] + [dn(b) for b in benefs] + ["拐力合计", "拐力占比"]):
            wsl.cell(rl, 2 + i, h)
        _style_header(wsl, rl, len(benefs) + 4)
        rl += 1
        givers = sorted(matrix.keys(), key=lambda g: -sum(matrix[g].values()))
        for g in givers:
            row_total = sum(matrix[g].values())
            wsl.cell(rl, 2, dn(g)).border = THIN
            for i, b in enumerate(benefs):
                v = matrix[g].get(b, 0.0)
                c = wsl.cell(rl, 3 + i, v if abs(v) > 0.5 else None)
                c.border = THIN
                c.number_format = "#,##0"
            c = wsl.cell(rl, 3 + len(benefs), row_total)
            c.font = BOLD
            c.number_format = "#,##0"
            c.border = THIN
            c = wsl.cell(rl, 4 + len(benefs), row_total / total if total else 0)
            c.number_format = "0.00%"
            c.border = THIN
            rl += 1
    else:
        wsl.cell(rl, 2, "本轴无干员间增益贡献").border = THIN
        rl += 1
    rl += 1

    # ---- 三、单击拐力 TOP15（拐在了哪一击上） ----
    wsl.cell(rl, 2, "三、单击拐力前 15（优化轴的方向：让大倍率命中吃满增益）").font = SECTION_FONT
    rl += 1
    for i, h in enumerate(["时间(s)", "拐人", "命中归属", "技能", "该击期望", "拐力", "占该击比例"]):
        wsl.cell(rl, 2 + i, h)
    _style_header(wsl, rl, 8)
    rl += 1
    for s, e, v in sorted(hit_contribs, key=lambda x: -x[2])[:15]:
        exp = e.get("exp") or 0
        vals = [e.get("t"), dn(s), dn(e.get("src") or ""),
                SKILL_TYPE_NAMES.get(e.get("st") or "", e.get("st") or "—"),
                exp, v, v / exp if exp else 0]
        for i, v2 in enumerate(vals):
            c = wsl.cell(rl, 2 + i, v2)
            c.border = THIN
            if i == 0:
                c.number_format = "0.00"
            elif i in (4, 5):
                c.number_format = "#,##0"
            elif i == 6:
                c.number_format = "0.0%"
        rl += 1

    # ---- 四、计算方法说明 ----
    rl += 1
    for line, is_head in LMDI_METHOD_LINES:
        c = wsl.cell(rl, 2, line)
        c.font = Font(size=9, bold=is_head,
                      color="1F3864" if is_head else "595959")
        rl += 1

    # ══════════════ Sheet 3.6 贡献归因 ══════════════
    # 全链路贡献归因 = LMDI 乘区拐力 + 技力催化 + 充能催化（附着催化单列经手展示）。
    wsp = wb.create_sheet("贡献归因")
    wsp.sheet_view.showGridLines = False
    wsp.column_dimensions["A"].width = 2
    for col, w in zip("BCDEFGHIJK", (12, 13, 12, 13, 12, 13, 14, 14, 10)):
        wsp.column_dimensions[col].width = w
    rp = 2
    for col in range(2, 12):
        wsp.cell(rp, col).fill = PatternFill("solid", fgColor=INK)
    wsp.cell(rp, 2, "贡献归因 · 全链路（LMDI拐力 + 技力催化 + 充能催化）").font = TITLE_FONT
    wsp.merge_cells(start_row=rp, start_column=2, end_row=rp, end_column=11)
    wsp.row_dimensions[rp].height = 26
    rp += 1
    hazard_bar(wsp, rp, 2, 11)
    rp += 2

    # ── 数据收集 ──
    sp_gen: dict = defaultdict(float)          # src → 技力产出（点）
    for h in hits:
        sp_gen[h.get("src")] += (h.get("spRecovery") or 0) + (h.get("spReturn") or 0)
    sp_cost_total = sum(-s["sp"] for s in sim.get("spLog", [])
                        if s.get("kind") == "SP_CHANGE" and (s.get("sp") or 0) < 0)
    dpsp_team = (total / sp_cost_total) if sp_cost_total else 0.0

    # 动作实例 → 所属干员 / 实例伤害（用于充能归因与终结技伤害折算）
    act_owner: dict = {}
    act_exp: dict = defaultdict(float)
    for h in sim.get("hits", []):
        act_owner.setdefault(h.get("act"), h.get("src"))
        act_exp[h.get("act")] += h.get("exp") or 0
    by_act = {k: [{"exp": v}] for k, v in act_exp.items()}
    charge_in: dict = defaultdict(float)       # (提供者, 接收者) → 充能点数
    for s2 in sim.get("spLog", []):
        if s2.get("kind") != "ULT_ENERGY_CHANGE" or (s2.get("gauge") or 0) <= 0:
            continue
        owner = act_owner.get(s2.get("source"))
        if owner and owner != s2.get("track"):
            charge_in[(owner, s2.get("track"))] += s2["gauge"]

    # 各 track 终结技：耗能（轴上的 gaugeCost）与实际伤害
    ult_cost: dict = defaultdict(float)
    ult_dmg: dict = defaultdict(float)
    sc = (project.get("scenarioList") or [{}])[0].get("data", {}) if isinstance(project, dict) else {}
    for t in sc.get("tracks", []):
        for a in t.get("actions", []):
            if a.get("skillKey") == "ultimate" or a.get("type") == "ultimate":
                cost = float(a.get("gaugeCost") or 0)
                if cost > 0:
                    ult_cost[t["id"]] += cost
                    ult_dmg[t["id"]] += sum(
                        hh.get("exp") or 0 for hh in by_act.get(a.get("instanceId"), []))

    # 附着施加者追踪：反应伤害的「附着催化（经手）」
    infl_last: dict = {}                        # element → (source, t)
    attach_cat: dict = defaultdict(float)       # src → 经手反应伤害
    enemy_events = sorted(sim.get("enemyLog", []),
                          key=lambda x: x.get("t") or 0)
    ev_idx = 0
    REACTION_KINDS = {"artsBurst", "corrosion", "combustion", "electrification",
                      "solidification", "shatter"}
    for h in sorted(damaging, key=lambda x: x["t"]):
        while ev_idx < len(enemy_events) and (enemy_events[ev_idx].get("t") or 0) <= h["t"]:
            e2 = enemy_events[ev_idx]
            if e2.get("type") == "INFLICTION_APPLY" and e2.get("element"):
                infl_last[e2["element"]] = (e2.get("source") or "", e2["t"])
            ev_idx += 1
        rx = h.get("reaction")
        if rx in REACTION_KINDS:
            src = infl_last.get(h.get("elem") or "", (None, None))[0]
            if src:
                attach_cat[src] += h.get("exp") or 0

    pulse_rows = []
    for tid in track_ids:
        row = lmdi.get(tid, {"dmg": 0.0, "buff": 0.0})
        g = sp_gen.get(tid, 0.0)
        sp_part = g * dpsp_team
        chg = sum(cnt * (ult_dmg.get(j, 0.0) / cost) for (prov, j), cnt in charge_in.items()
                  for cost in [ult_cost.get(j, 0.0)] if cost > 0 and prov == tid)
        pulse_rows.append({
            "track": tid, "self": row.get("dmg", 0.0), "buff": row.get("buff", 0.0),
            "spGen": g, "spPart": sp_part,
            "charge": sum(v for (p, j), v in charge_in.items() if p == tid),
            "chgPart": chg,
            "attach": attach_cat.get(tid, 0.0),
        })
    pulse_total_sum = sum(r["buff"] + r["spPart"] + r["chgPart"] for r in pulse_rows)

    for i, h in enumerate(["干员", "LMDI自身结算", "LMDI拐力", "技力产出(点)", "技力催化伤害",
                           "充能提供(点)", "充能催化伤害", "附着催化(经手)", "贡献归因合计", "占比"]):
        wsp.cell(rp, 2 + i, h)
    _style_header(wsp, rp, 11)
    rp += 1
    for r in sorted(pulse_rows, key=lambda r: -(r["buff"] + r["spPart"] + r["chgPart"])):
        total_i = r["buff"] + r["spPart"] + r["chgPart"]
        vals = [dn(r["track"]), r["self"], r["buff"], r["spGen"], r["spPart"],
                r["charge"], r["chgPart"], r["attach"], total_i,
                (total_i / total if total else 0)]
        for i, v in enumerate(vals):
            c = wsp.cell(rp, 2 + i, v)
            c.border = THIN
            c.font = data_font
            if i in (1, 2, 4, 6, 7, 8):
                c.number_format = "#,##0"
            elif i in (3, 5):
                c.number_format = "#,##0.0"
            elif i == 9:
                c.number_format = "0.0%"
        rp += 1
    c = wsp.cell(rp, 2, "全队合计")
    c.font = BOLD
    vals = ["—",
            sum(r["self"] for r in pulse_rows), sum(r["buff"] for r in pulse_rows),
            sum(r["spGen"] for r in pulse_rows), sum(r["spPart"] for r in pulse_rows),
            sum(r["charge"] for r in pulse_rows), sum(r["chgPart"] for r in pulse_rows),
            sum(r["attach"] for r in pulse_rows), pulse_total_sum,
            (pulse_total_sum / total if total else 0)]
    for i, v in enumerate(vals):
        c = wsp.cell(rp, 2 + i, v)
        c.font = BOLD
        c.border = THIN
        if i in (1, 2, 4, 6, 7, 8):
            c.number_format = "#,##0"
        elif i in (3, 5):
            c.number_format = "#,##0.0"
        elif i == 9:
            c.number_format = "0.0%"
    rp += 2

    for line, is_head in [
        ("全链路贡献归因＝LMDI 乘区拐力之上，"
         "把「资源供给」按机会成本折算成伤害的三段扩展。", True),
        ("① 技力催化：干员产出的每 1 点技力，按全队平均「每点技力伤害」（总伤害 ÷ 全队技力总消耗）折算。"
         "先锋的技力供给由此显性入账。", False),
        ("② 充能催化：为他人提供的每 1 点终结技能量，按「受益者终结技实际伤害 ÷ 其耗能」折算。"
         "辅助的充能天赋（如洁尔佩塔类）由此入账。", False),
        ("③ 附着催化（经手，单列）：其施加的附着被消耗引发的反应伤害。该伤害已计入触发者的直伤，"
         "故不重复计入 贡献归因合计，仅作经手展示。", False),
        ("贡献归因合计 ＝ LMDI 拐力 + 技力催化 + 充能催化，全队精确可加；"
         "自身结算与直伤仍以「拐力归因（LMDI）」表为准。", True),
    ]:
        c = wsp.cell(rp, 2, line)
        c.font = Font(size=9, bold=is_head, color="1F3864" if is_head else "595959")
        rp += 1

    # ══════════════ Sheet 3.5 伤害时间轴 ══════════════
    wst = wb.create_sheet("伤害时间轴")
    wst.sheet_view.showGridLines = False
    wst.column_dimensions["A"].width = 2
    for col, w in zip("BCDEF", (10, 15, 11, 12, 11)):
        wst.column_dimensions[col].width = w
    wst.column_dimensions["G"].width = 34
    rt = 2
    for col in range(2, 8):
        wst.cell(rt, col).fill = PatternFill("solid", fgColor=INK)
    wst.cell(rt, 2, "伤害时间轴 · 每秒输出 / 爆发与空转 / 失衡窗口").font = TITLE_FONT
    wst.merge_cells(start_row=rt, start_column=2, end_row=rt, end_column=7)
    wst.row_dimensions[rt].height = 26
    rt += 1
    hazard_bar(wst, rt, 2, 7)
    rt += 2

    per_sec: dict = defaultdict(lambda: {"dmg": 0.0, "n": 0})
    for e in damaging:
        d = per_sec[int(e["t"])]
        d["dmg"] += e["exp"]
        d["n"] += 1

    # 失衡窗口（由逐击 的 broken 状态翻转推得）
    windows: list[tuple[float, float]] = []
    cur_start = None
    last_t = 0.0
    for e in sorted(damaging, key=lambda x: x["t"]):
        last_t = e["t"]
        b = bool((e.get("enemyState") or {}).get("broken"))
        if b and cur_start is None:
            cur_start = e["t"]
        elif not b and cur_start is not None:
            windows.append((cur_start, e["t"]))
            cur_start = None
    if cur_start is not None:
        windows.append((cur_start, last_t))

    def _in_window(t: float) -> bool:
        return any(a <= t <= b for a, b in windows)

    if per_sec:
        t0, t1 = min(per_sec), max(per_sec)

        # 一、每秒伤害 + 黄色数据条
        wst.cell(rt, 2, "一、每秒伤害（数据条＝输出强度；灰底行＝处于失衡窗口）").font = SECTION_FONT
        rt += 1
        for i, h in enumerate(["秒", "伤害", "占比", "累计占比", "命中数"]):
            wst.cell(rt, 2 + i, h)
        _style_header(wst, rt, 6)
        rt += 1
        first_data = rt
        stagger_fill = PatternFill("solid", fgColor="EDEDED")
        cum = 0.0
        for sec in range(t0, t1 + 1):
            d = per_sec.get(sec, {"dmg": 0.0, "n": 0})
            cum += d["dmg"]
            vals = [sec, d["dmg"], d["dmg"] / total if total else 0,
                    cum / total if total else 0, d["n"]]
            in_stag = _in_window(sec + 0.5)
            for i, v in enumerate(vals):
                c = wst.cell(rt, 2 + i, v)
                c.border = THIN
                c.font = data_font
                if i == 0:
                    c.number_format = "0"
                elif i == 1:
                    c.number_format = "#,##0"
                elif i in (2, 3):
                    c.number_format = "0.0%"
                if in_stag:
                    c.fill = stagger_fill
            rt += 1
        wst.conditional_formatting.add(
            f"C{first_data}:C{rt - 1}",
            DataBarRule(start_type="num", start_value=0, end_type="max",
                        color=YELLOW, showValue=True),
        )
        rt += 1

        # 二、失衡窗口利用
        wst.cell(rt, 2, "二、失衡窗口利用（失衡期是倍率最高的输出窗口，务必吃满）").font = SECTION_FONT
        rt += 1
        for i, h in enumerate(["开始(s)", "结束(s)", "时长(s)", "窗口内伤害", "占全部伤害"]):
            wst.cell(rt, 2 + i, h)
        _style_header(wst, rt, 6)
        rt += 1
        for a, b_ in windows:
            dmg_in = sum(e["exp"] for e in damaging if a <= e["t"] <= b_)
            vals = [a, b_, b_ - a, dmg_in, dmg_in / total if total else 0]
            for i, v in enumerate(vals):
                c = wst.cell(rt, 2 + i, v)
                c.border = THIN
                c.font = data_font
                if i in (0, 1, 2):
                    c.number_format = "0.00"
                elif i == 3:
                    c.number_format = "#,##0"
                elif i == 4:
                    c.number_format = "0.0%"
            rt += 1
        rt += 1

        # 三、爆发 TOP5 秒
        wst.cell(rt, 2, "三、爆发最高的 5 秒（检查增益是否恰好覆盖这几秒）").font = SECTION_FONT
        rt += 1
        for i, h in enumerate(["秒", "伤害", "命中数", "该秒主要动作"]):
            wst.cell(rt, 2 + i, h)
        _style_header(wst, rt, 5, underline=GR_DMG)
        rt += 1
        for sec, d in sorted(per_sec.items(), key=lambda kv: -kv[1]["dmg"])[:5]:
            acts: dict = defaultdict(float)
            for e in damaging:
                if int(e["t"]) == sec:
                    acts[e.get("actionName") or e.get("st") or "?"] += e["exp"]
            top_act = max(acts.items(), key=lambda kv: kv[1])[0] if acts else "—"
            vals = [sec, d["dmg"], d["n"], top_act]
            for i, v in enumerate(vals):
                c = wst.cell(rt, 2 + i, v)
                c.border = THIN
                c.font = data_font
                if i == 1:
                    c.number_format = "#,##0"
            rt += 1

    # ══════════════ Sheet 3 技能与元素分布 ══════════════
    ws3 = wb.create_sheet("技能与元素分布")
    ws3.sheet_view.showGridLines = False
    for col, w in zip("BCDEFG", (16, 14, 12, 16, 14, 12)):
        ws3.column_dimensions[col].width = w
    r3 = 2
    ws3.cell(r3, 2, "按技能类型").font = SECTION_FONT
    r3 += 1
    by_type = defaultdict(float)
    cnt_type = Counter()
    for e in damaging:
        by_type[e.get("st") or "?"] += e["exp"]
        cnt_type[e.get("st") or "?"] += 1
    for i, h in enumerate(["类型", "伤害", "占比", "命中数"]):
        ws3.cell(r3, 2 + i, h)
    _style_header(ws3, r3, 5)
    r3 += 1
    for k, v in sorted(by_type.items(), key=lambda kv: -kv[1]):
        ws3.cell(r3, 2, SKILL_TYPE_NAMES.get(k, k)).border = THIN
        c = ws3.cell(r3, 3, v); c.number_format = "#,##0"; c.border = THIN
        c = ws3.cell(r3, 4, v / total if total else 0); c.number_format = "0.00%"; c.border = THIN
        ws3.cell(r3, 5, cnt_type[k]).border = THIN
        r3 += 1
    r3 += 1
    ws3.cell(r3, 2, "按元素").font = SECTION_FONT
    r3 += 1
    by_elem = defaultdict(float)
    for e in damaging:
        by_elem[e.get("elem") or "?"] += e["exp"]
    for i, h in enumerate(["元素", "伤害", "占比"]):
        ws3.cell(r3, 2 + i, h)
    _style_header(ws3, r3, 4)
    r3 += 1
    for k, v in sorted(by_elem.items(), key=lambda kv: -kv[1]):
        ws3.cell(r3, 2, ELEMENT_NAMES.get(k, k)).border = THIN
        c = ws3.cell(r3, 3, v); c.number_format = "#,##0"; c.border = THIN
        c = ws3.cell(r3, 4, v / total if total else 0); c.number_format = "0.00%"; c.border = THIN
        r3 += 1
    r3 += 1
    ws3.cell(r3, 2, "干员 × 技能类型").font = SECTION_FONT
    r3 += 1
    types = [k for k, _ in sorted(by_type.items(), key=lambda kv: -kv[1])]
    ws3.cell(r3, 2, "干员")
    for i, k in enumerate(types):
        ws3.cell(r3, 3 + i, SKILL_TYPE_NAMES.get(k, k))
    ws3.cell(r3, 3 + len(types), "合计")
    _style_header(ws3, r3, 3 + len(types))
    r3 += 1
    dealt_by_track = defaultdict(float)
    for e in damaging:
        dealt_by_track[e.get("src")] += e["exp"]
    for tid in track_ids:
        ws3.cell(r3, 2, dn(tid)).border = THIN
        for i, k in enumerate(types):
            v = sum(e["exp"] for e in damaging if e.get("src") == tid and e.get("st") == k)
            c = ws3.cell(r3, 3 + i, v)
            c.number_format = "#,##0"; c.border = THIN
        c = ws3.cell(r3, 3 + len(types), dealt_by_track.get(tid, 0.0))
        c.number_format = "#,##0"; c.font = BOLD; c.border = THIN
        r3 += 1
    r3 += 1

    # 三、按干员 × 具体动作（深度玩家视角：每个技能各打了多少）
    ws3.cell(r3, 2, "三、按干员 × 具体动作").font = SECTION_FONT
    r3 += 1
    for i, h in enumerate(["干员", "动作", "命中数", "伤害", "均伤", "最高单击", "占比"]):
        ws3.cell(r3, 2 + i, h)
    _style_header(ws3, r3, 8)
    r3 += 1
    by_action: dict = defaultdict(lambda: {"n": 0, "dmg": 0.0, "max": 0.0})
    for e in damaging:
        key = (e.get("src") or "?", e.get("actionName") or SKILL_TYPE_NAMES.get(e.get("st") or "", e.get("st") or "?"))
        d = by_action[key]
        d["n"] += 1
        d["dmg"] += e["exp"]
        d["max"] = max(d["max"], e["exp"])
    action_rows = sorted(by_action.items(), key=lambda kv: -kv[1]["dmg"])
    for (src, act), d in action_rows:
        vals = [dn(src), act, d["n"], d["dmg"], d["dmg"] / d["n"] if d["n"] else 0,
                d["max"], d["dmg"] / total if total else 0]
        for i, v in enumerate(vals):
            c = ws3.cell(r3, 2 + i, v)
            c.border = THIN
            if i in (3, 4, 5):
                c.number_format = "#,##0"
            elif i == 6:
                c.number_format = "0.00%"
        r3 += 1

    # ══════════════ Sheet 4 SP与能量 ══════════════
    # 动作实例ID → 可读动作名（spLog 的 source 只带 inst_xxx）
    act_labels: dict = {}
    for h in damaging:
        act = h.get("act")
        if act and act not in act_labels:
            act_labels[act] = hit_label(h)
    ws4 = wb.create_sheet("SP与能量")
    ws4.sheet_view.showGridLines = False
    for i, w in enumerate((9, 12, 14, 10, 10, 18)):
        ws4.column_dimensions[get_column_letter(i + 1)].width = w
    r4 = 2
    for col in range(1, 7):
        ws4.cell(r4, col).fill = PatternFill("solid", fgColor=INK)
    ws4.cell(r4, 1, "SP 与能量 · 技力池曲线与流水").font = TITLE_FONT
    ws4.merge_cells(start_row=r4, start_column=1, end_row=r4, end_column=6)
    ws4.row_dimensions[r4].height = 26
    r4 += 1
    hazard_bar(ws4, r4, 1, 6)
    r4 += 2

    sp_series = sim.get("spSeries") or []
    sp_rate = sim.get("spRegenRate")
    max_sp = sim.get("maxSp") or 300
    axis_end = summary.get("lastDamageTime", 0)

    def _sp_at(t: float):
        """SP 序列是事件折线，线性插值取任意时刻的存量。"""
        if not sp_series:
            return None
        if t <= sp_series[0]["time"]:
            return sp_series[0]["sp"]
        for a, b in zip(sp_series, sp_series[1:]):
            if a["time"] <= t <= b["time"]:
                span = b["time"] - a["time"]
                if span <= 1e-9:
                    return b["sp"]
                return a["sp"] + (b["sp"] - a["sp"]) * (t - a["time"]) / span
        return sp_series[-1]["sp"]

    # 一、SP 池概览（引擎官方投影：含自然恢复/上限/暂停/欠费）
    ws4.cell(r4, 1, "一、SP 池概览（引擎官方投影，含自然恢复）").font = SECTION_FONT
    r4 += 1
    if sp_series:
        in_axis = [p for p in sp_series if p.get("time", 0) <= axis_end + 1e-6] or sp_series
        sp_vals = [p.get("sp", 0) for p in in_axis]
        debt_peak = max(p.get("debtSp", 0) for p in in_axis)
        full_time = next((p["time"] for p in sp_series
                          if p["time"] > axis_end and p.get("sp", 0) >= max_sp - 1e-6), None)
        verdict = ("SP 全程健康，未触底" if min(sp_vals) > 1
                   else "SP 一度触底（贴近卡费），注意充能与费用安排")
        if debt_peak > 0.5:
            verdict += f"；曾欠费 {debt_peak:.0f} 点"
        for i, h in enumerate(["初始", "窗口最低", "窗口最高", "窗口结束", "自然恢复"]):
            ws4.cell(r4, 1 + i, h)
        _style_header(ws4, r4, 5)
        r4 += 1
        for i, v in enumerate([in_axis[0].get("sp", 0), min(sp_vals), max(sp_vals),
                               in_axis[-1].get("sp", 0),
                               f"{sp_rate:g} /s" if sp_rate else "—"]):
            c = ws4.cell(r4, 1 + i, v)
            c.border = THIN
            c.font = data_font
            if i < 4:
                c.number_format = "#,##0.0"
        r4 += 1
        for lab, val in (("循环评价", verdict),
                         ("SP 回满时刻",
                          f"{full_time:.1f}s（轴结束后 {full_time - axis_end:.1f}s）" if full_time else "—")):
            ws4.cell(r4, 1, lab).font = BOLD
            c = ws4.cell(r4, 2, val)
            c.border = THIN
            ws4.merge_cells(start_row=r4, start_column=2, end_row=r4, end_column=5)
            r4 += 1
        r4 += 1

        # 二、SP 曲线（按秒 + 数据条）
        ws4.cell(r4, 1, "二、SP 曲线（按秒，数据条＝技力存量，蓝色）").font = SECTION_FONT
        r4 += 1
        for i, h in enumerate(["秒", "SP 存量", "占上限"]):
            ws4.cell(r4, 1 + i, h)
        _style_header(ws4, r4, 3, underline="7FA8D9")
        r4 += 1
        first_sp_row = r4
        for sec in range(int(sp_series[0]["time"]), int(axis_end) + 1):
            v = _sp_at(sec)
            c1 = ws4.cell(r4, 1, sec); c1.border = THIN; c1.font = data_font
            c2 = ws4.cell(r4, 2, v); c2.border = THIN; c2.font = data_font
            c2.number_format = "#,##0.0"
            c3 = ws4.cell(r4, 3, (v or 0) / max_sp); c3.border = THIN; c3.font = data_font
            c3.number_format = "0%"
            r4 += 1
        ws4.conditional_formatting.add(
            f"B{first_sp_row}:B{r4 - 1}",
            DataBarRule(start_type="num", start_value=0, end_type="num",
                        end_value=max_sp, color="5B9BD5", showValue=True),
        )
        r4 += 1
    else:
        r4 += 1

    # 三、事件流水
    ws4.cell(r4, 1, "三、事件流水（技劘支出 / 回复 / 终结技能量）").font = SECTION_FONT
    r4 += 1
    for i, h in enumerate(["时间(s)", "干员", "事件", "SP变动", "能量变动", "来源"]):
        ws4.cell(r4, 1 + i, h)
    _style_header(ws4, r4, 6)
    r4 += 1
    for s in sorted(sim.get("spLog", []), key=lambda x: x["t"]):
        kind = s.get("kind") or ""
        sp = s.get("sp") if kind == "SP_CHANGE" else None
        gauge = s.get("gauge") if kind == "ULT_ENERGY_CHANGE" else None
        event = "技力" if kind == "SP_CHANGE" else "终结技能量"
        reason = SP_REASON_NAMES.get(s.get("reason") or "", s.get("reason") or "")
        source = s.get("source") or ""
        # source 是动作实例ID（inst_xxx）时换成「干员+动作」；其余（如终结技能量的
        # 实例ID）对用户没有意义，留空——干员列已经说明是谁
        source = act_labels.get(source, "") if source else ""
        vals = [s["t"], dn(s.get("track") or ""), event, sp, gauge,
                (f"{reason} · {source}" if reason and source
                 else (source or reason or ""))]
        for i, v in enumerate(vals):
            c = ws4.cell(r4, 1 + i, v)
            c.border = THIN
            c.font = data_font
            if i == 0:
                c.number_format = "0.00"
            elif i in (3, 4):
                c.number_format = "#,##0.0"
        r4 += 1

    # ══════════════ Sheet 5 增益时间线 ══════════════
    ws6 = wb.create_sheet("增益时间线")
    for i, h in enumerate(["时间(s)", "对象", "增益/减益", "类别", "数值", "层数", "来源", "失效(s)"]):
        ws6.cell(1, i + 1, h)
    _style_header(ws6, 1, 8)
    for i, w in enumerate((9, 10, 26, 10, 10, 8, 10, 9)):
        ws6.column_dimensions[get_column_letter(i + 1)].width = w
    ws6.freeze_panes = "A2"
    z6 = 2
    for e in sorted(sim.get("operatorLog", []), key=lambda x: x["t"]):
        stat = e.get("stat") or ""
        vals = [e["t"], dn(e.get("target") or ""), T(e.get("id") or ""),
                STAT_NAMES.get(stat, T(stat)) if stat else "状态", e.get("value"), e.get("cumulative") or e.get("stacks"),
                dn(e.get("source") or ""),
                f"{e['expires']:.1f}" if e.get("expires") is not None else "永久"]
        for i, v in enumerate(vals):
            c = ws6.cell(z6, i + 1, v)
            c.border = THIN
            if i == 0 or i == 7:
                c.number_format = "0.00"
            elif i == 4:
                c.number_format = "0.##"
        z6 += 1
    for e in sorted(sim.get("enemyLog", []), key=lambda x: x["t"]):
        etype = e.get("type") or ""
        if etype == "INFLICTION_APPLY":
            vals = [e["t"], "敌人", T(f"{e.get('element') or ''}Infliction"), "附着",
                    "—", e.get("stacks"), dn(e.get("source") or ""), "—"]
        elif etype == "DEBUFF_APPLY" or (e.get("id") == "corrosion:resShred" and etype == "ENEMY_EFFECT_APPLY"):
            vals = [e["t"], "敌人", T(e.get("id") or "corrosion"), "减益",
                    e.get("value"), e.get("stacks"), dn(e.get("source") or ""), "—"]
        elif etype == "CORROSION_TICK":
            vals = [e["t"], "敌人", "腐蚀", "减益", e.get("value"), "—",
                    dn(e.get("source") or ""), "—"]
        else:
            continue
        for i, v in enumerate(vals):
            c = ws6.cell(z6, i + 1, v)
            c.border = THIN
            if i == 0:
                c.number_format = "0.00"
            elif i == 4:
                c.number_format = "0.##"
        z6 += 1

    # ══════════════ Sheet 6 增益覆盖统计 ══════════════
    ws5 = wb.create_sheet("增益覆盖统计")
    ws5.sheet_view.showGridLines = False
    for i, w in enumerate((30, 12, 14, 13, 12, 18, 15)):
        ws5.column_dimensions[get_column_letter(i + 1)].width = w
    r5 = 2
    for col in range(1, 8):
        ws5.cell(r5, col).fill = PatternFill("solid", fgColor=INK)
    ws5.cell(r5, 1, "增益覆盖统计 · 时间覆盖率与伤害加权覆盖率").font = TITLE_FONT
    ws5.merge_cells(start_row=r5, start_column=1, end_row=r5, end_column=7)
    ws5.row_dimensions[r5].height = 26
    r5 += 1
    hazard_bar(ws5, r5, 1, 7)
    r5 += 2
    note = ("时间覆盖率＝增益生效时长 ÷ 战斗窗口；伤害加权覆盖率＝"
            "对象在该增益生效期间打出的伤害 ÷ 其总伤害——"
            "后者才反映增益对输出的实际意义。")
    c = ws5.cell(r5, 1, note)
    c.font = Font(size=9, color="595959")
    ws5.merge_cells(start_row=r5, start_column=1, end_row=r5, end_column=7)
    r5 += 2

    axis_start = summary.get("startlineSec", 0)
    axis_end = summary.get("lastDamageTime", 0)
    span = max(1e-6, axis_end - axis_start)

    # 按 (增益, 对象) 合并生效区间
    intervals: dict = defaultdict(list)
    meta: dict = {}
    for e in sim.get("operatorLog", []):
        key = (e.get("id") or "?", e.get("target") or "?")
        start = float(e.get("t") or 0)
        end = float(e["expires"]) if e.get("expires") is not None else axis_end
        end = min(max(end, start), axis_end)
        intervals[key].append((start, max(start, end)))
        meta[key] = {"source": e.get("source") or ""}

    dmg_by_src: dict = defaultdict(float)
    for e in damaging:
        dmg_by_src[e.get("src")] += e["exp"]

    cov_rows = []
    for (bid, target), ivs in intervals.items():
        clipped = sorted((max(a, axis_start), min(b, axis_end))
                         for a, b in ivs if b > axis_start)
        merged: list[list[float]] = []
        for a, b in clipped:
            if merged and a <= merged[-1][1] + 1e-9:
                merged[-1][1] = max(merged[-1][1], b)
            else:
                merged.append([a, b])
        duration = sum(b - a for a, b in merged)
        if duration <= 1e-6:
            continue
        dmg_in = sum(e["exp"] for e in damaging
                     if e.get("src") == target
                     and any(a <= e["t"] <= b for a, b in merged))
        own = dmg_by_src.get(target, 0.0)
        cov_rows.append({
            "id": bid, "target": target,
            "source": meta.get((bid, target), {}).get("source", ""),
            "duration": duration,
            "timeCov": duration / span,
            "dmgIn": dmg_in,
            "dmgCov": (dmg_in / own) if own > 0 else 0.0,
        })
    cov_rows.sort(key=lambda r: -r["dmgIn"])

    for i, h in enumerate(["增益/减益", "对象", "来源", "覆盖时长(s)",
                           "时间覆盖率", "生效期伤害(对象)", "伤害加权覆盖率"]):
        ws5.cell(r5, 1 + i, h)
    _style_header(ws5, r5, 7)
    r5 += 1
    for r in cov_rows:
        vals = [T(r["id"]), dn(r["target"]), dn(r["source"]) or "—",
                r["duration"], r["timeCov"], r["dmgIn"], r["dmgCov"]]
        for i, v in enumerate(vals):
            c = ws5.cell(r5, 1 + i, v)
            c.border = THIN
            c.font = data_font
            if i == 3:
                c.number_format = "0.00"
            elif i in (4, 6):
                c.number_format = "0.0%"
            elif i == 5:
                c.number_format = "#,##0"
        r5 += 1

    return wb


def sp_by_track(track_id: str, sim: dict) -> float:
    return sum(-s["sp"] for s in sim.get("spLog", [])
               if s.get("kind") == "SP_CHANGE" and s.get("track") == track_id
               and (s.get("sp") or 0) < 0)


def write_report(sim: dict, project: dict, out_path: str, names: dict | None = None) -> str:
    wb = build_workbook(sim, project, names)
    wb.save(out_path)
    return out_path
